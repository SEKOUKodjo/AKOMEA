# -*- coding: utf-8 -*-
"""Genere la presentation PowerPoint (<= 10 diapositives)."""
import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BASE = Path(__file__).parent.parent
DATA = BASE / "data"
FIG = Path(__file__).parent / "figures"
OUT = Path(__file__).parent / "Presentation_Economie_Numerique_Togo.pptx"
META = json.loads((DATA / "meta.json").read_text())

GREEN = RGBColor(0x0B, 0x6E, 0x4F)
GREEN2 = RGBColor(0x00, 0x87, 0x51)
YELLOW = RGBColor(0xFF, 0xCE, 0x00)
RED = RGBColor(0xD2, 0x10, 0x34)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1C, 0x2B, 0x27)
GREY = RGBColor(0x6b, 0x7b, 0x76)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def nf(v, dec=0):
    return (f"{v:,.{dec}f}").replace(",", " ").replace(".", ",")


def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, color, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def txt(s, x, y, w, h, text, size=18, color=INK, bold=False, italic=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri"):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color; r.font.name = font
    return tb


def bullets(s, x, y, w, h, items, size=16, color=INK):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = "•  " + it
        r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = "Calibri"
        p.space_after = Pt(8)
    return tb


def flagbar(s, y=Emu(0), h=Inches(0.13)):
    w = SW
    rect(s, 0, y, int(w * 0.4), h, GREEN)
    rect(s, int(w * 0.4), y, int(w * 0.22), h, YELLOW)
    rect(s, int(w * 0.62), y, int(w * 0.38), h, RED)


def header(s, title, num):
    rect(s, 0, 0, SW, Inches(1.0), GREEN)
    flagbar(s, Inches(1.0))
    txt(s, Inches(0.4), Inches(0.12), Inches(11.5), Inches(0.8), title,
        size=26, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(12.2), Inches(0.12), Inches(0.9), Inches(0.8), f"{num}/10",
        size=14, color=YELLOW, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)


def kpi(s, x, y, w, value, label, color, tcolor=WHITE):
    rect(s, x, y, w, Inches(1.25), color)
    txt(s, x, y + Inches(0.12), w, Inches(0.6), value, size=30, color=tcolor,
        bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x, y + Inches(0.72), w, Inches(0.5), label, size=12, color=tcolor,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def placeholder(s, x, y, w, h, label):
    from pptx.enum.shapes import MSO_SHAPE
    sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor(0xFB, 0xFD, 0xFC)
    sp.line.color.rgb = GREEN2; sp.line.width = Pt(1.5)
    sp.line.dash_style = None
    sp.shadow.inherit = False
    tf = sp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = "📷 " + label
    r.font.size = Pt(14); r.font.color.rgb = GREY; r.font.italic = True
    return sp


def pic(s, name, x, y, w=None, h=None):
    if (FIG / name).exists():
        s.shapes.add_picture(str(FIG / name), x, y, width=w, height=h)


# ===================== SLIDE 1 — TITRE =====================
s = slide()
rect(s, 0, 0, SW, SH, WHITE)
rect(s, 0, 0, SW, Inches(2.1), GREEN)
flagbar(s, Inches(2.1), Inches(0.16))
txt(s, Inches(0.5), Inches(0.35), Inches(12.3), Inches(0.5),
    "REPUBLIQUE TOGOLAISE  —  Togo AI Lab", size=16, color=YELLOW, bold=True,
    align=PP_ALIGN.CENTER)
txt(s, Inches(0.5), Inches(0.9), Inches(12.3), Inches(1.1),
    "Acces aux telecommunications & services numeriques au Togo",
    size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
txt(s, Inches(0.5), Inches(2.5), Inches(12.3), Inches(0.7),
    "Diagnostic territorial · Zones blanches · Priorisation & investissements",
    size=18, color=INK, bold=True, align=PP_ALIGN.CENTER)
kpi(s, Inches(0.9), Inches(3.6), Inches(2.6), nf(META["n_mobile_money"]),
    "Agents mobile money", GREEN)
kpi(s, Inches(3.75), Inches(3.6), Inches(2.6), str(META["n_agences_telecom"]),
    "Agences telecoms", GREEN2)
kpi(s, Inches(6.6), Inches(3.6), Inches(2.6), str(META["n_datacenters"]),
    "Datacenters", GREY)
kpi(s, Inches(9.45), Inches(3.6), Inches(2.9), str(META["n_zones_blanches"]),
    "Zones blanches (cantons)", RED)
txt(s, Inches(0.5), Inches(5.4), Inches(12.3), Inches(0.6),
    f"Population de reference : projection 2026 = {nf(META['national_pop_2026'])} hab. "
    f"(RGPH-5 2022, taux {META['taux_croissance']*100:.1f} %/an)",
    size=14, color=GREY, italic=True, align=PP_ALIGN.CENTER)
txt(s, Inches(0.5), Inches(6.5), Inches(12.3), Inches(0.6),
    "Maurice Kodjo SEKOU — Analyste Statisticien  |  Shiny for Python · Highcharts",
    size=13, color=GREEN, bold=True, align=PP_ALIGN.CENTER)

# ===================== SLIDE 2 — CONTEXTE =====================
s = slide(); header(s, "Contexte & objectifs", 2)
txt(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(0.9),
    "A partir des donnees ouvertes sur les infrastructures telecoms, les services "
    "numeriques et la demographie, dresser un diagnostic territorial de l'acces au "
    "numerique et prioriser l'extension de la connectivite et de l'inclusion.",
    size=16, color=INK)
bullets(s, Inches(0.7), Inches(2.6), Inches(12), Inches(4), [
    "Cartographier les agences des operateurs (Togocom, Moov, CANAL+) et les datacenters.",
    "Analyser la couverture des services (mobile money) et son adequation a la population.",
    "Croiser la localisation des infrastructures avec la densite demographique.",
    "Identifier les zones blanches numeriques (localites non desservies).",
    "Proposer des recommandations ciblees et chiffrees (FCFA / EUR).",
], size=17)

# ===================== SLIDE 3 — DONNEES & METHODE =====================
s = slide(); header(s, "Donnees & methode", 3)
txt(s, Inches(0.5), Inches(1.3), Inches(6), Inches(0.4), "Donnees mobilisees",
    size=18, color=RED, bold=True)
bullets(s, Inches(0.6), Inches(1.8), Inches(6.2), Inches(4), [
    f"Agences telecoms : {META['n_agences_telecom']} (Togocom 62, Moov 28).",
    f"Agents mobile money : {nf(META['n_mobile_money'])} points geolocalises.",
    f"Datacenters : {META['n_datacenters']} etablissements.",
    "Population RGPH-5 2022 (pays → canton).",
    "Limites administratives geoBoundaries (adm1/2/3).",
], size=15)
txt(s, Inches(6.9), Inches(1.3), Inches(6), Inches(0.4), "Traitements",
    size=18, color=RED, bold=True)
bullets(s, Inches(7.0), Inches(1.8), Inches(6), Inches(4), [
    "Nettoyage : coordonnees, libelles normalises, annees.",
    "Reconstruction de la hierarchie de population (validee par les sommes).",
    "Projection demographique 2026 (taux 2,3 %/an).",
    "Indicateurs de densite, concentration, correlation.",
    "Pre-agregation en .parquet → chargement rapide.",
], size=15)

# ===================== SLIDE 4 — OFFRE & REPARTITION =====================
s = slide(); header(s, "L'offre numerique et sa repartition", 4)
pic(s, "fig_mm_region.png", Inches(0.4), Inches(1.5), w=Inches(6.6))
pic(s, "fig_operateurs.png", Inches(7.4), Inches(1.5), h=Inches(3.6))
txt(s, Inches(0.5), Inches(5.5), Inches(12.3), Inches(1.6),
    f"Densite nationale : {nf(META['mm_pour_1000_national'],2)} agents mobile money "
    f"pour 1 000 hab. et {nf(META['agences_pour_100k_national'],2)} agence pour "
    "100 000 hab. La region Maritime domine ; Togocom porte l'essentiel du reseau "
    "d'agences. Les 3 datacenters sont tous situes dans le Grand Lome.",
    size=15, color=INK)

# ===================== SLIDE 5 — INEGALITES =====================
s = slide(); header(s, "De fortes inegalites territoriales", 5)
pic(s, "fig_densite.png", Inches(0.35), Inches(1.5), w=Inches(7.6))
rect(s, Inches(8.2), Inches(1.6), Inches(4.8), Inches(4.6), RGBColor(0xF3, 0xF9, 0xF6))
txt(s, Inches(8.4), Inches(1.75), Inches(4.5), Inches(0.5), "Chiffres cles",
    size=17, color=GREEN, bold=True)
bullets(s, Inches(8.4), Inches(2.35), Inches(4.5), Inches(3.8), [
    "Densite de 0,42 a 5,36 /1000 hab.",
    "Rapport de 1 a 12,8 entre extremes.",
    "Moyenne 1,98 · mediane 1,52.",
    "Coefficient de variation : 60,6 %.",
    "Correlation pop / agents : r = 0,90.",
], size=15, color=INK)
txt(s, Inches(0.5), Inches(6.6), Inches(12), Inches(0.6),
    "L'offre suit la demographie mais laisse subsister d'importants sous-equipements relatifs.",
    size=13, color=GREY, italic=True)

# ===================== SLIDE 6 — CONCENTRATION & ZONES BLANCHES =====================
s = slide(); header(s, "Concentration & zones blanches", 6)
pic(s, "fig_zones_blanches.png", Inches(0.4), Inches(1.5), w=Inches(6.8))
kpi(s, Inches(7.6), Inches(1.7), Inches(2.5), "25,9 %", "Agents dans le Golfe", RED)
kpi(s, Inches(10.3), Inches(1.7), Inches(2.5), "54,4 %", "Agents top-5 prefectures", GREY)
kpi(s, Inches(7.6), Inches(3.2), Inches(2.5), str(META["n_zones_blanches"]),
    "Cantons sans agent", RED)
kpi(s, Inches(10.3), Inches(3.2), Inches(2.5),
    f"{nf((META['n_cantons']-META['n_zones_blanches'])/META['n_cantons']*100,1)} %",
    "Taux de couverture", GREEN)
bullets(s, Inches(7.6), Inches(4.8), Inches(5.4), Inches(2), [
    "12 prefectures sur 39 sans agence telecom.",
    "Datacenters 100 % concentres a Lome.",
    "Zones blanches : Plateaux, Centrale, Kara.",
], size=14)

# ===================== SLIDE 7 — PRIORITES =====================
s = slide(); header(s, "Priorisation territoriale", 7)
pic(s, "fig_priorites.png", Inches(0.4), Inches(1.5), w=Inches(7.2))
txt(s, Inches(7.9), Inches(1.55), Inches(5.1), Inches(0.9),
    "Score de priorite (0-100)", size=16, color=RED, bold=True)
txt(s, Inches(7.9), Inches(2.25), Inches(5.1), Inches(1.6),
    "Indicateur composite : deficit de densite mobile money (45 %), deficit de "
    "densite d'agences (35 %) et poids demographique (20 %).", size=14, color=INK)
placeholder(s, Inches(7.9), Inches(3.9), Inches(5.1), Inches(2.7),
            "Capture : carte des priorites\n(onglet « Carte » du dashboard)")

# ===================== SLIDE 8 — INVESTISSEMENT =====================
s = slide(); header(s, "Estimation des investissements", 8)
txt(s, Inches(0.5), Inches(1.35), Inches(12.3), Inches(0.9),
    "Simulateur interactif : pour un territoire et un objectif de densite choisis, "
    "le tableau de bord calcule les points a creer et le budget en FCFA et en EUR "
    "(parite 1 € = 655,957 FCFA).", size=15, color=INK)
txt(s, Inches(0.5), Inches(2.5), Inches(12.3), Inches(0.5),
    "Chiffrage national — objectif : 3 agents/1000 hab. & 1,5 agence/100 000 hab.",
    size=15, color=GREEN, bold=True)
kpi(s, Inches(0.9), Inches(3.1), Inches(3.6), "+6 800", "Points mobile money", GREEN)
kpi(s, Inches(4.8), Inches(3.1), Inches(3.6), "+43", "Agences telecoms", GREEN2)
kpi(s, Inches(8.7), Inches(3.1), Inches(3.7), "78", "Zones blanches a resorber", RED)
rect(s, Inches(0.9), Inches(4.7), Inches(5.7), Inches(1.5), GREEN)
txt(s, Inches(0.9), Inches(4.85), Inches(5.7), Inches(0.5), "Investissement total",
    size=15, color=WHITE, align=PP_ALIGN.CENTER)
txt(s, Inches(0.9), Inches(5.35), Inches(5.7), Inches(0.7), "3 303 500 000 FCFA",
    size=26, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
rect(s, Inches(6.8), Inches(4.7), Inches(5.6), Inches(1.5), RED)
txt(s, Inches(6.8), Inches(4.85), Inches(5.6), Inches(0.5), "Equivalent euros",
    size=15, color=WHITE, align=PP_ALIGN.CENTER)
txt(s, Inches(6.8), Inches(5.35), Inches(5.6), Inches(0.7), "≈ 5 036 153 €",
    size=26, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(0.5), Inches(6.4), Inches(12.3), Inches(0.6),
    "Budget entierement modulable (objectifs et couts unitaires) et ventilable par prefecture.",
    size=13, color=GREY, italic=True, align=PP_ALIGN.CENTER)

# ===================== SLIDE 9 — RECOMMANDATIONS =====================
s = slide(); header(s, "Recommandations strategiques", 9)
bullets(s, Inches(0.7), Inches(1.5), Inches(12), Inches(5), [
    "Cibler les prefectures a score eleve : Kpendjal, Mo, Kpendjal-Ouest, Tandjoare, Est-Mono.",
    "Resorber les 78 zones blanches via des reseaux d'agents mobile money de proximite.",
    "Deconcentrer les datacenters hors du Grand Lome (resilience & souverainete des donnees).",
    "Ouvrir des agences telecoms dans les 12 prefectures qui en sont depourvues.",
    "Piloter les investissements par le simulateur (objectifs de densite, budget FCFA/EUR par territoire).",
], size=18)

# ===================== SLIDE 10 — CONCLUSION =====================
s = slide(); header(s, "Tableau de bord & conclusion", 10)
placeholder(s, Inches(0.5), Inches(1.5), Inches(7.2), Inches(4.6),
            "Capture d'ecran du tableau de bord\n(page d'accueil — a inserer)")
txt(s, Inches(8.0), Inches(1.6), Inches(5.0), Inches(0.5), "Ce que livre l'outil",
    size=17, color=RED, bold=True)
bullets(s, Inches(8.0), Inches(2.2), Inches(5.0), Inches(3.8), [
    "Diagnostic clair et territorialise.",
    "8 onglets interactifs, filtres region/prefecture.",
    "Cartes Highmaps & graphiques Highcharts.",
    "Simulateur d'investissement FCFA/EUR.",
    "Application autonome, chargement rapide.",
], size=15)
rect(s, 0, Inches(6.6), SW, Inches(0.9), GREEN)
txt(s, Inches(0.5), Inches(6.6), Inches(12.3), Inches(0.9),
    "Maurice Kodjo SEKOU — Analyste Statisticien  ·  Data Challenge Economie Numerique, Togo AI Lab",
    size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

prs.save(str(OUT))
print("Presentation ecrite :", OUT.name, "|", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
