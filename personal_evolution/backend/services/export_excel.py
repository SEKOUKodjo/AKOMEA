"""Export Excel des données et des indicateurs de PEI.

Le classeur contient les données brutes (une feuille par type de donnée) et des
indicateurs calculés par des FORMULES Excel qui pointent vers ces données : si
l'utilisateur corrige une valeur dans Excel, les indicateurs se mettent à jour.
Seules les analyses statistiques (tests, corrélations, recommandations) sont des
valeurs calculées par PEI au moment de l'export ; elles sont signalées comme telles.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date, datetime, timedelta

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from .. import models as m
from ..i18n import tr, weekday as weekday_name
from ..domain import (DIMENSIONS, GOAL_STATUSES, HORIZONS, INTENTION_STATUSES, REFLECTION_KINDS, label as label_of,
                      dimension_label)
from . import analytics
from . import skills as skill_service

GREEN, YELLOW, RED = "006A4E", "FFCE00", "D21034"
FONT = "Arial"
F_BODY = Font(name=FONT, size=10)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=16, bold=True, color=GREEN)
F_SUB = Font(name=FONT, size=10, italic=True, color="5B6A63")
F_SECTION = Font(name=FONT, size=12, bold=True, color=GREEN)
F_BOLD = Font(name=FONT, size=10, bold=True)
FILL_HEAD = PatternFill("solid", fgColor=GREEN)
FILL_KPI = PatternFill("solid", fgColor="F4F6F2")
BORDER_HEAD = Border(bottom=Side(style="medium", color=YELLOW))
BORDER_ROW = Border(bottom=Side(style="thin", color="DFE6E0"))

FMT_DATE = "dd/mm/yyyy"
FMT_PCT = "0.0%"
FMT_DEC = "0.0"
FMT_INT = "0"

DIM_LABELS = [v["label"] for v in DIMENSIONS.values()]


@dataclass
class Col:
    key: str
    header: str
    width: int = 14
    fmt: str | None = None
    wrap: bool = False


class Table:
    """Feuille de données : en-têtes stylés, repérage des colonnes par clé pour écrire des formules."""

    def __init__(self, ws: Worksheet, cols: list[Col], header_row: int = 1):
        self.ws, self.cols, self.header_row = ws, cols, header_row
        self.letter = {c.key: get_column_letter(i + 1) for i, c in enumerate(cols)}
        self.n = 0
        for i, c in enumerate(cols, start=1):
            cell = ws.cell(row=header_row, column=i, value=c.header)
            cell.font, cell.fill, cell.border = F_HEAD, FILL_HEAD, BORDER_HEAD
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = c.width
        ws.row_dimensions[header_row].height = 30
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    @property
    def first(self) -> int:
        return self.header_row + 1

    @property
    def last(self) -> int:
        return self.header_row + max(self.n, 1)

    def add(self, values: dict) -> int:
        r = self.first + self.n
        for i, c in enumerate(self.cols, start=1):
            v = values.get(c.key)
            if callable(v):
                v = v(r)
            cell = self.ws.cell(row=r, column=i, value=v)
            cell.font = F_BODY
            cell.border = BORDER_ROW
            if c.fmt:
                cell.number_format = c.fmt
            if c.wrap:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        self.n += 1
        return r

    def rng(self, key: str, absolute: bool = True) -> str:
        """Plage d'une colonne, qualifiée par le nom de la feuille."""
        col = self.letter[key]
        d = "$" if absolute else ""
        return f"'{self.ws.title}'!{d}{col}{d}{self.first}:{d}{col}{d}{self.last}"

    def finish(self) -> None:
        last_col = get_column_letter(len(self.cols))
        self.ws.auto_filter.ref = f"A{self.header_row}:{last_col}{self.header_row + max(self.n, 1)}"


def _title(ws: Worksheet, title: str, subtitle: str) -> None:
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws["A2"] = subtitle
    ws["A2"].font = F_SUB


def _note(ws: Worksheet, row: int, text: str, col: int = 1) -> None:
    c = ws.cell(row=row, column=col, value=text)
    c.font = F_SUB


def _clean(v):
    if v is None:
        return None
    if isinstance(v, float) and pd.isna(v):
        return None
    if hasattr(v, "item"):
        return v.item()
    return v


def build_workbook(session: Session, start: date | None = None, end: date | None = None) -> Workbook:
    fr = analytics.load_frames(session, start, end)
    daily = analytics.daily_series(fr, start, end) if (start and end) else analytics.daily_series(fr)
    if daily.empty:
        today = pd.Timestamp(date.today())
        daily = analytics.daily_series(fr, today.date(), today.date())
    p_start, p_end = daily.index.min().date(), daily.index.max().date()

    # Libellés utilisés à la fois dans les données et dans les formules : même langue partout.
    S_PREVU, S_REAL, S_PART = tr("Prévu"), tr("Réalisé"), tr("Partiel")
    S_NR, S_AB, S_REP = tr("Non réalisé"), tr("Abandonné"), tr("Reporté")
    YES, NO = tr("Oui"), tr("Non")
    E_ON, E_LATE, E_MISSED = (tr(analytics.ENTRY_STATUS[k]) for k in ("a_temps", "rattrapee", "oubliee"))
    WD = ",".join(f'"{weekday_name(i)}"' for i in range(7))

    wb = Workbook()
    ws_syn = wb.active
    ws_syn.title = tr("Synthese")
    ws_q = wb.create_sheet(tr("Quotidien"))
    ws_w = wb.create_sheet(tr("Hebdomadaire"))
    ws_mo = wb.create_sheet(tr("Mensuel"))
    ws_i = wb.create_sheet(tr("Intentions"))
    ws_a = wb.create_sheet(tr("Activites"))
    ws_r = wb.create_sheet(tr("Reflexions"))
    ws_cal = wb.create_sheet(tr("Calibration"))
    ws_ab = wb.create_sheet(tr("Abandons"))
    ws_g = wb.create_sheet(tr("Objectifs"))
    ws_s = wb.create_sheet(tr("Competences"))
    ws_e = wb.create_sheet(tr("Evaluations"))
    ws_an = wb.create_sheet(tr("Analyses"))
    ws_d = wb.create_sheet(tr("Decisions"))
    ws_x = wb.create_sheet(tr("Experiences"))
    ws_p = wb.create_sheet(tr("Relations"))
    ws_m = wb.create_sheet(tr("Medias"))

    # Intentions (données brutes)
    goals = {g.id: g.title for g in session.query(m.Goal).all()}
    ti = Table(ws_i, [
        Col("id", tr("N°"), 7), Col("date", tr("Date"), 12, FMT_DATE), Col("dim", tr("Dimension"), 18), Col("cat", tr("Type d'activité"), 20),
        Col("desc", tr("Intention"), 42, wrap=True), Col("prio", tr("Priorité"), 10), Col("est", tr("Durée prévue (min)"), 12, FMT_INT),
        Col("status", tr("Statut"), 13), Col("act", tr("Durée réelle (min)"), 12, FMT_INT), Col("gap", tr("Écart (min)"), 11, FMT_INT),
        Col("score", tr("Score de réalisation"), 12, "0.0"), Col("result", tr("Résultat"), 34, wrap=True),
        Col("reason", tr("Raison d'un écart"), 26, wrap=True), Col("goal", tr("Objectif relié"), 30, wrap=True),
    ])
    prio = {1: tr("Haute"), 2: tr("Normale"), 3: tr("Basse")}
    it = fr.intentions.sort_values(["date", "id"]) if not fr.intentions.empty else fr.intentions
    for row in it.itertuples():
        ti.add({
            "id": int(row.id), "date": row.date.date(), "dim": dimension_label(row.dimension), "cat": row.category,
            "desc": row.description, "prio": prio.get(int(row.priority), ""), "est": _clean(row.estimated_minutes),
            "status": label_of(INTENTION_STATUSES, row.status), "act": _clean(row.actual_minutes),
            "gap": lambda r: f'=IF(AND(ISNUMBER({ti.letter["est"]}{r}),ISNUMBER({ti.letter["act"]}{r})),'
                             f'{ti.letter["act"]}{r}-{ti.letter["est"]}{r},"")',
            "score": lambda r: (f'=IF({ti.letter["status"]}{r}="{S_REAL}",1,IF({ti.letter["status"]}{r}="{S_PART}",0.5,'
                                f'IF({ti.letter["status"]}{r}="{S_PREVU}","",0)))'),
            "result": _clean(row.result), "reason": _clean(row.reason),
            "goal": goals.get(int(row.goal_id)) if _clean(row.goal_id) else None,
        })
    ti.finish()
    ws_i.cell(row=1, column=ti.cols.index(next(c for c in ti.cols if c.key == "score")) + 1).comment = Comment(
        tr("Réalisé = 1, Partiel = 0,5, Non réalisé, Abandonné ou Reporté = 0, Prévu = non évalué (vide)."), "PEI")

    # Activités (données brutes)
    skills = {s.id: s.name for s in session.query(m.Skill).all()}
    ta = Table(ws_a, [
        Col("id", tr("N°"), 7), Col("date", tr("Date"), 12, FMT_DATE), Col("dim", tr("Dimension"), 18), Col("cat", tr("Type d'activité"), 20),
        Col("desc", tr("Activité"), 42, wrap=True), Col("min", tr("Durée (min)"), 11, FMT_INT), Col("start", tr("Début"), 8),
        Col("end", tr("Fin"), 8), Col("result", tr("Résultat"), 36, wrap=True), Col("planned", tr("Prévue"), 9),
        Col("skill", tr("Compétence"), 18),
    ])
    acts = {a.id: a for a in session.query(m.Activity).filter(m.Activity.id.in_(list(fr.activities["id"]))).all()} \
        if not fr.activities.empty else {}
    ac = fr.activities.sort_values(["date", "id"]) if not fr.activities.empty else fr.activities
    for row in ac.itertuples():
        a = acts.get(int(row.id))
        ta.add({
            "id": int(row.id), "date": row.date.date(), "dim": dimension_label(row.dimension), "cat": row.category,
            "desc": row.description, "min": int(row.minutes), "start": a.start_time if a else None,
            "end": a.end_time if a else None, "result": _clean(row.result),
            "planned": YES if _clean(row.intention_id) else NO,
            "skill": skills.get(int(row.skill_id)) if _clean(row.skill_id) else None,
        })
    ta.finish()

    # Réflexions
    t_ref = Table(ws_r, [Col("date", tr("Date"), 12, FMT_DATE), Col("kind", tr("Type"), 18), Col("dim", tr("Dimension"), 18),
                      Col("content", tr("Contenu"), 80, wrap=True)])
    rf = fr.reflections.sort_values("date") if not fr.reflections.empty else fr.reflections
    for row in rf.itertuples():
        t_ref.add({"date": row.date.date(), "kind": label_of(REFLECTION_KINDS, row.kind),
                "dim": dimension_label(row.dimension) if _clean(row.dimension) else None, "content": row.content})
    t_ref.finish()

    # Quotidien : une ligne par jour, indicateurs par formules
    highlights = {}
    if not fr.days.empty:
        highlights = dict(zip(fr.days["date"], fr.days["highlight"]))
    qcols = [
        Col("date", tr("Date"), 12, FMT_DATE), Col("wd", tr("Jour"), 11), Col("logged", tr("Renseigné"), 10),
        Col("n", tr("Intentions"), 10, FMT_INT), Col("eval", tr("Évaluées"), 10, FMT_INT), Col("done", tr("Réalisées"), 10, FMT_INT),
        Col("rate", tr("Taux de réalisation"), 12, FMT_PCT), Col("planned", tr("Temps prévu (min)"), 12, FMT_INT),
        Col("actual", tr("Temps réalisé (min)"), 12, FMT_INT),
    ] + [Col(f"d_{k}", f"{tr(v['label'])} (min)", 12, FMT_INT) for k, v in DIMENSIONS.items()] + [
        Col("sleep_hours", tr("Sommeil (h)"), 10, FMT_DEC), Col("sleep_quality", tr("Qualité sommeil"), 10, FMT_INT),
        Col("energy", tr("Énergie"), 9, FMT_INT), Col("motivation", tr("Motivation"), 10, FMT_INT), Col("mood", tr("Humeur"), 9, FMT_INT),
        Col("stress", tr("Stress"), 9, FMT_INT), Col("screen_minutes", tr("Écran (min)"), 10, FMT_INT),
        Col("steps", tr("Pas"), 10, "#,##0"), Col("water_liters", tr("Eau (l)"), 9, FMT_DEC),
        Col("n_ref", tr("Réflexions"), 10, FMT_INT), Col("entry", tr("Saisie"), 22), Col("entry_delay", tr("Délai de saisie (jours)"), 11, FMT_INT),
        Col("highlight", tr("Moment marquant"), 40, wrap=True),
    ]
    tq = Table(ws_q, qcols)
    I, A, R = ti, ta, t_ref

    def q_formulas(r: int) -> dict:
        d = f"$A{r}"
        out = {
            "wd": f'=INDEX({{{WD}}},WEEKDAY({d},2))',
            "n": f"=COUNTIFS({I.rng('date')},{d})",
            "eval": f'=COUNTIFS({I.rng("date")},{d},{I.rng("status")},"<>{S_PREVU}")',
            "done": (f'=COUNTIFS({I.rng("date")},{d},{I.rng("status")},"{S_REAL}")'
                     f'+COUNTIFS({I.rng("date")},{d},{I.rng("status")},"{S_PART}")'),
            "rate": f'=IF({tq.letter["eval"]}{r}=0,"",SUMIFS({I.rng("score")},{I.rng("date")},{d})/{tq.letter["eval"]}{r})',
            "planned": f"=SUMIFS({I.rng('est')},{I.rng('date')},{d})",
            "actual": f"=SUMIFS({A.rng('min')},{A.rng('date')},{d})",
            "n_ref": f"=COUNTIFS({R.rng('date')},{d})",
        }
        for k, v in DIMENSIONS.items():
            out[f"d_{k}"] = f'=SUMIFS({A.rng("min")},{A.rng("date")},{d},{A.rng("dim")},"{tr(v["label"])}")'
        return out

    entry = analytics.entry_discipline(session, p_start, p_end)
    entry_by_day = {e["date"]: e for e in entry["days"]}
    for ts, row in daily.iterrows():
        r = tq.first + tq.n
        vals = {"date": ts.date(), "logged": YES if row["logged"] else NO, **q_formulas(r)}
        for k in ["sleep_hours", "sleep_quality", "energy", "motivation", "mood", "stress", "screen_minutes", "steps",
                  "water_liters"]:
            vals[k] = _clean(row.get(k))
        vals["highlight"] = _clean(highlights.get(ts))
        ent = entry_by_day.get(str(ts.date()))
        if ent:
            vals["entry"], vals["entry_delay"] = ent["status_label"], ent["delay"]
        tq.add(vals)
    tq.finish()
    Q = tq

    # Hebdomadaire et mensuel : agrégats par formules
    def period_table(ws: Worksheet, label_end: str, starts: list[date], end_formula) -> Table:
        t = Table(ws, [
            Col("start", tr("Début"), 12, FMT_DATE), Col("end", label_end, 12, FMT_DATE),
            Col("logged", tr("Jours renseignés"), 11, FMT_INT), Col("reg", tr("Régularité"), 11, FMT_PCT),
            Col("n", tr("Intentions"), 10, FMT_INT), Col("eval", tr("Évaluées"), 10, FMT_INT), Col("done", tr("Réalisées"), 10, FMT_INT),
            Col("rate", tr("Taux de réalisation"), 12, FMT_PCT), Col("hours", tr("Heures suivies"), 11, FMT_DEC),
        ] + [Col(f"h_{k}", f"{tr(v['label'])} (h)", 11, FMT_DEC) for k, v in DIMENSIONS.items()] + [
            Col("sleep", tr("Sommeil moyen (h)"), 11, FMT_DEC), Col("energy", tr("Énergie moyenne"), 11, "0.00"),
            Col("mood", tr("Humeur moyenne"), 11, "0.00"), Col("gap", tr("Écart moyen d'estimation (min)"), 14, FMT_DEC),
        ])
        for s in starts:
            r = t.first + t.n
            a, b = f"$A{r}", f"$B{r}"
            inq = f'{Q.rng("date")},">="&{a},{Q.rng("date")},"<="&{b}'
            ini = f'{I.rng("date")},">="&{a},{I.rng("date")},"<="&{b}'
            vals = {
                "start": s, "end": end_formula(r),
                "logged": f'=COUNTIFS({inq},{Q.rng("logged")},"{YES}")',
                "reg": f'=IFERROR(C{r}/COUNTIFS({inq}),"")',
                "n": f"=SUMIFS({Q.rng('n')},{inq})", "eval": f"=SUMIFS({Q.rng('eval')},{inq})",
                "done": f"=SUMIFS({Q.rng('done')},{inq})",
                "rate": f'=IF({t.letter["eval"]}{r}=0,"",SUMIFS({I.rng("score")},{ini})/{t.letter["eval"]}{r})',
                "hours": f"=SUMIFS({Q.rng('actual')},{inq})/60",
                "sleep": f'=IFERROR(AVERAGEIFS({Q.rng("sleep_hours")},{inq}),"")',
                "energy": f'=IFERROR(AVERAGEIFS({Q.rng("energy")},{inq}),"")',
                "mood": f'=IFERROR(AVERAGEIFS({Q.rng("mood")},{inq}),"")',
                "gap": f'=IFERROR(AVERAGEIFS({I.rng("gap")},{ini}),"")',
            }
            for k in DIMENSIONS:
                vals[f"h_{k}"] = f"=SUMIFS({Q.rng('d_' + k)},{inq})/60"
            t.add(vals)
        t.finish()
        return t

    w0 = p_start - timedelta(days=p_start.weekday())
    weeks = [w0 + timedelta(days=7 * k) for k in range((p_end - w0).days // 7 + 1)]
    tw = period_table(ws_w, tr("Fin"), weeks, lambda r: f"=A{r}+6")
    months, mcur = [], p_start.replace(day=1)
    while mcur <= p_end:
        months.append(mcur)
        mcur = (mcur.replace(day=28) + timedelta(days=4)).replace(day=1)
    period_table(ws_mo, tr("Fin"), months, lambda r: f"=EOMONTH(A{r},0)")

    # Calibration du temps : formules par type d'activité
    tc = Table(ws_cal, [
        Col("cat", tr("Type d'activité"), 24), Col("n", tr("Observations"), 12, FMT_INT), Col("est", tr("Prévu moyen (min)"), 13, FMT_DEC),
        Col("act", tr("Réel moyen (min)"), 13, FMT_DEC), Col("gap", tr("Écart moyen (min)"), 13, FMT_DEC),
        Col("ratio", tr("Réel / prévu"), 11, "0.00"), Col("read", tr("Lecture"), 26),
    ], header_row=4)
    _title(ws_cal, tr("Calibration du temps"), tr("Temps prévu contre temps réellement utilisé, par type d'activité "
                                                  "(intentions avec durée prévue et réelle)."))
    cats = sorted(set(it["category"].dropna())) if not it.empty else []
    for cat in cats:
        tc.add({
            "cat": cat,
            "n": lambda r: f'=COUNTIFS({I.rng("cat")},$A{r},{I.rng("est")},">0",{I.rng("act")},">0")',
            "est": lambda r: f'=IFERROR(AVERAGEIFS({I.rng("est")},{I.rng("cat")},$A{r},{I.rng("est")},">0",{I.rng("act")},">0"),"")',
            "act": lambda r: f'=IFERROR(AVERAGEIFS({I.rng("act")},{I.rng("cat")},$A{r},{I.rng("est")},">0",{I.rng("act")},">0"),"")',
            "gap": lambda r: f'=IFERROR(D{r}-C{r},"")',
            "ratio": lambda r: f'=IFERROR(D{r}/C{r},"")',
            "read": lambda r: (f'=IF(B{r}<3,"{tr("Pas assez de données")}",IF(F{r}>1.1,"{tr("Sous estimation")}",'
                               f'IF(F{r}<0.9,"{tr("Surestimation")}","{tr("Estimation juste")}")))'),
        })
    tc.finish()
    _note(ws_cal, tc.last + 2, tr("Seuils : rapport réel / prévu supérieur à 1,1 = sous estimation, inférieur à 0,9 = "
                                  "surestimation. Au moins 3 observations sont nécessaires."))

    # Abandons : formules par type d'activité, puis raisons citées
    _title(ws_ab, tr("Objectifs non réalisés"), tr("Part des intentions évaluées qui n'ont pas été réalisées, par type d'activité."))
    tb = Table(ws_ab, [
        Col("cat", tr("Type d'activité"), 24), Col("eval", tr("Évaluées"), 10, FMT_INT), Col("nr", tr("Non réalisées"), 12, FMT_INT),
        Col("ab", tr("Abandonnées"), 12, FMT_INT), Col("rep", tr("Reportées"), 10, FMT_INT), Col("rate", tr("Taux de non réalisation"), 14, FMT_PCT),
    ], header_row=4)
    for cat in cats:
        tb.add({
            "cat": cat,
            "eval": lambda r: f'=COUNTIFS({I.rng("cat")},$A{r},{I.rng("status")},"<>{S_PREVU}")',
            "nr": lambda r: f'=COUNTIFS({I.rng("cat")},$A{r},{I.rng("status")},"{S_NR}")',
            "ab": lambda r: f'=COUNTIFS({I.rng("cat")},$A{r},{I.rng("status")},"{S_AB}")',
            "rep": lambda r: f'=COUNTIFS({I.rng("cat")},$A{r},{I.rng("status")},"{S_REP}")',
            "rate": lambda r: f'=IF(B{r}=0,"",(C{r}+D{r})/B{r})',
        })
    tb.finish()
    reasons = sorted(set(it["reason"].dropna().str.strip())) if not it.empty else []
    start_row = tb.last + 3
    ws_ab.cell(row=start_row - 1, column=1, value=tr("Raisons citées")).font = F_SECTION
    tr2 = Table(ws_ab, [Col("reason", tr("Raison"), 24), Col("count", tr("Nombre de fois"), 10, FMT_INT)], header_row=start_row)
    ws_ab.freeze_panes = "A5"
    for reason in reasons:
        tr2.add({"reason": reason, "count": lambda r: f'=COUNTIFS({I.rng("reason")},$A{r})'})

    # Objectifs (valeurs calculées par PEI, sous objectifs inclus)
    _title(ws_g, tr("Objectifs"), tr("Progression calculée par PEI au moment de l'export, en incluant les actions des sous objectifs."))
    tg = Table(ws_g, [
        Col("title", tr("Objectif"), 40, wrap=True), Col("horizon", tr("Horizon"), 13), Col("dim", tr("Dimension"), 18),
        Col("status", tr("Statut"), 11), Col("parent", tr("Objectif parent"), 32, wrap=True), Col("target", tr("Échéance"), 12, FMT_DATE),
        Col("n", tr("Actions reliées"), 11, FMT_INT), Col("rate", tr("Taux de réalisation"), 12, FMT_PCT),
        Col("hours", tr("Heures investies"), 11, FMT_DEC), Col("last", tr("Dernière action"), 12, FMT_DATE),
    ], header_row=4)
    for g in analytics.goals_progress(session):
        tg.add({"title": g["title"], "horizon": label_of(HORIZONS, g["horizon"]),
                "dim": dimension_label(g["dimension"]) if g["dimension"] else None,
                "status": label_of(GOAL_STATUSES, g["status"]),
                "parent": goals.get(g["parent_id"]) if g["parent_id"] else None,
                "target": date.fromisoformat(g["target_date"]) if g["target_date"] else None,
                "n": g["n_intentions"], "rate": g["completion_rate"], "hours": g["hours"],
                "last": date.fromisoformat(g["last_action"]) if g["last_action"] else None})
    tg.finish()

    # Compétences et évaluations
    te = Table(ws_e, [Col("skill", tr("Compétence"), 22), Col("date", tr("Date"), 12, FMT_DATE), Col("rating", tr("Auto évaluation (sur 5)"), 13, "0.0"),
                      Col("score", tr("Score objectif (sur 100)"), 13, FMT_INT), Col("etype", tr("Type de preuve"), 14),
                      Col("evidence", tr("Preuve"), 40, wrap=True), Col("comment", tr("Commentaire"), 30, wrap=True)])
    skill_objs = session.query(m.Skill).order_by(m.Skill.name).all()
    for s in skill_objs:
        for p in s.progress:
            te.add({"skill": s.name, "date": p.date, "rating": p.self_rating, "score": p.objective_score,
                    "etype": p.evidence_type, "evidence": p.evidence, "comment": p.comment})
    te.finish()
    _title(ws_s, tr("Compétences"), tr("Ressenti (dernière auto évaluation) et preuves observables (heures et séances sur la période)."))
    tsk = Table(ws_s, [
        Col("name", tr("Compétence"), 24), Col("target", tr("Niveau visé"), 11, FMT_INT), Col("rating", tr("Dernière auto évaluation"), 13, "0.0"),
        Col("first", tr("Première auto évaluation"), 13, "0.0"), Col("change", tr("Évolution du ressenti"), 12, "+0.0;-0.0;0"),
        Col("hours", tr("Heures de pratique"), 12, FMT_DEC), Col("sessions", tr("Séances"), 10, FMT_INT),
        Col("results", tr("Séances avec résultat"), 12, FMT_INT), Col("evals", tr("Évaluations"), 11, FMT_INT),
    ], header_row=4)
    for s in skill_objs:
        summ = skill_service.summary(session, s)
        tsk.add({
            "name": s.name, "target": s.target_level, "rating": summ["current_rating"], "first": summ["first_rating"],
            "change": lambda r: f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r})),C{r}-D{r},"")',
            "hours": lambda r: f"=SUMIFS({A.rng('min')},{A.rng('skill')},$A{r})/60",
            "sessions": lambda r: f"=COUNTIFS({A.rng('skill')},$A{r})",
            "results": lambda r: f'=COUNTIFS({A.rng("skill")},$A{r},{A.rng("result")},"<>")',
            "evals": lambda r: f"=COUNTIFS({te.rng('skill')},$A{r})",
        })
    tsk.finish()

    # Analyses statistiques (valeurs calculées par PEI)
    _title(ws_an, tr("Analyses statistiques"), tr("Valeurs calculées par PEI au moment de l'export.") + " "
                                              + analytics.causality_note())
    row = 4

    def block(title: str, method: str, cols: list[Col], rows: list[dict]) -> None:
        nonlocal row
        ws_an.cell(row=row, column=1, value=title).font = F_SECTION
        _note(ws_an, row + 1, method)
        t = Table(ws_an, cols, header_row=row + 2)
        for rr in rows:
            t.add(rr)
        if not rows:
            _note(ws_an, row + 3, tr("Pas assez de données pour cette analyse."))
        row = t.last + 3

    trends = analytics.trends(daily, fr=fr)
    block(tr("Tendances des 3 dernières semaines"), tr("Comparaison avec les 3 semaines précédentes, test de Mann Whitney (p)."),
          [Col("label", tr("Indicateur"), 34), Col("dir", tr("Sens"), 10), Col("before", tr("Avant"), 10, "0.0"),
           Col("recent", tr("Récent"), 10, "0.0"), Col("p", tr("p"), 10, "0.0000"), Col("msg", tr("Lecture"), 70, wrap=True)],
          [{"label": t["label"], "dir": t["direction_label"], "before": t["before"], "recent": t["recent"],
            "p": t["p_value"], "msg": t["message"]} for t in trends])
    corr = analytics.correlations(daily)
    block(tr("Associations entre indicateurs"), tr("Corrélation de Spearman (rho), q = p valeur corrigée de Benjamini Hochberg."),
          [Col("x", tr("Variable 1"), 34), Col("y", tr("Variable 2"), 22), Col("rho", tr("rho"), 10, "0.000"), Col("n", tr("Jours"), 10, FMT_INT),
           Col("q", tr("q"), 10, "0.0000"), Col("read", tr("Lecture"), 30)],
          [{"x": p["x_label"], "y": p["y_label"], "rho": p["rho"], "n": p["n"], "q": p["q_value"],
            "read": f"{p['strength']}, {tr('nette') if p['significant'] else tr('incertaine')}"} for p in corr["pairs"][:20]])
    block(tr("Jours atypiques"), tr("Écart à la médiane glissante sur 28 jours, mesuré en écarts robustes (MAD)."),
          [Col("date", tr("Date"), 34, FMT_DATE), Col("label", tr("Indicateur"), 22), Col("value", tr("Valeur"), 10, "0.0"),
           Col("exp", tr("Habituel"), 10, "0.0"), Col("z", tr("Écart robuste"), 10, "0.0"), Col("dir", tr("Sens"), 30)],
          [{"date": date.fromisoformat(a["date"]), "label": a["label"], "value": a["value"], "exp": a["expected"],
            "z": a["z"], "dir": a["direction_label"]} for a in analytics.anomalies(daily)[:30]])
    eff = analytics.intention_count_effect(fr, daily)
    block(tr("Nombre d'intentions et réalisation"), tr("Taux de réalisation moyen selon le nombre d'intentions fixées le matin."),
          [Col("bucket", tr("Intentions fixées"), 34), Col("rate", tr("Taux moyen"), 22, FMT_PCT), Col("n", tr("Jours"), 10, FMT_INT)],
          [{"bucket": b["bucket"], "rate": b["completion_rate"], "n": b["n_days"]} for b in (eff or {}).get("buckets", [])])
    block(tr("Recommandations"), tr("Générées à partir des analyses ci dessus, avec leur mise en garde."),
          [Col("title", tr("Recommandation"), 34), Col("msg", tr("Message"), 70, wrap=True), Col("caveat", tr("Mise en garde"), 50, wrap=True)],
          [{"title": i["title"], "msg": i["message"], "caveat": i.get("caveat")} for i in analytics.insights(session, start, end)])
    ws_an.freeze_panes = None

    # Décisions, expériences, relations, médias
    td = Table(ws_d, [Col("date", tr("Date"), 12, FMT_DATE), Col("title", tr("Décision"), 34, wrap=True), Col("context", tr("Contexte"), 30, wrap=True),
                      Col("reasons", tr("Raisons"), 30, wrap=True), Col("alt", tr("Alternatives"), 24, wrap=True), Col("choice", tr("Choix"), 24, wrap=True),
                      Col("conf", tr("Confiance (sur 5)"), 11, FMT_INT), Col("review", tr("Revue prévue"), 12, FMT_DATE),
                      Col("outcome", tr("Résultat"), 30, wrap=True), Col("rating", tr("Appréciation (sur 5)"), 12, FMT_INT),
                      Col("lessons", tr("Leçon"), 30, wrap=True)])
    for d in session.query(m.Decision).order_by(m.Decision.date).all():
        td.add({"date": d.date, "title": d.title, "context": d.context, "reasons": d.reasons, "alt": d.alternatives,
                "choice": d.choice, "conf": d.confidence, "review": d.review_date, "outcome": d.outcome,
                "rating": d.outcome_rating, "lessons": d.lessons})
    td.finish()
    tx = Table(ws_x, [Col("title", tr("Expérience"), 28, wrap=True), Col("hyp", tr("Hypothèse"), 40, wrap=True), Col("start", tr("Début"), 12, FMT_DATE),
                      Col("end", tr("Fin"), 12, FMT_DATE), Col("metric", tr("Indicateur"), 20), Col("base", tr("Moyenne avant"), 11, "0.00"),
                      Col("exp", tr("Moyenne pendant"), 11, "0.00"), Col("p", tr("p"), 9, "0.000"), Col("msg", tr("Lecture"), 50, wrap=True)])
    for e in session.query(m.Experiment).order_by(m.Experiment.start_date).all():
        an = analytics.experiment_analysis(session, e)
        tx.add({"title": e.title, "hyp": e.hypothesis, "start": e.start_date, "end": e.end_date, "metric": an.get("label"),
                "base": an.get("baseline_mean"), "exp": an.get("experiment_mean"), "p": an.get("p_value"), "msg": an.get("message")})
    tx.finish()
    tp = Table(ws_p, [Col("name", tr("Personne"), 22), Col("rel", tr("Relation"), 14), Col("every", tr("Contact souhaité (jours)"), 13, FMT_INT),
                      Col("n", tr("Échanges"), 10, FMT_INT), Col("last", tr("Dernier échange"), 13, FMT_DATE),
                      Col("since", tr("Jours depuis"), 11, FMT_INT), Col("todo", tr("À relancer"), 11)])
    for p in session.query(m.Person).order_by(m.Person.name).all():
        last = max((i.date for i in p.interactions), default=None)
        tp.add({"name": p.name, "rel": p.relation, "every": p.contact_every_days, "n": len(p.interactions), "last": last,
                "since": lambda r: f'=IF(ISNUMBER(E{r}),TODAY()-E{r},"")',
                "todo": lambda r: f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(F{r})),IF(F{r}>C{r},"{YES}","{NO}"),"")'})
    tp.finish()
    tmd = Table(ws_m, [Col("date", tr("Date"), 12, FMT_DATE), Col("type", tr("Type"), 10), Col("name", tr("Fichier"), 30), Col("caption", tr("Légende"), 30, wrap=True),
                       Col("size", tr("Taille (Ko)"), 11, "#,##0"), Col("status", tr("Transcription"), 13),
                       Col("text", tr("Texte transcrit"), 60, wrap=True), Col("sha", tr("Empreinte SHA 256"), 22)])
    mq = session.query(m.Media).outerjoin(m.Day)
    if start:
        mq = mq.filter(m.Day.date >= start)
    if end:
        mq = mq.filter(m.Day.date <= end)
    for md in mq.order_by(m.Media.created_at).all():
        tmd.add({"date": md.day.date if md.day else md.created_at.date(), "type": md.type, "name": md.original_name,
                 "caption": md.caption, "size": round((md.size_bytes or 0) / 1024, 1),
                 "status": md.transcription_status, "text": md.transcription, "sha": md.sha256})
    tmd.finish()

    # Synthèse : indicateurs clés par formules
    ws = ws_syn
    now = datetime.now()
    _title(ws, "Personal Evolution Intelligence", tr("Export du {d} à {t}, période du {a} au {b}", d=f"{now:%d/%m/%Y}",
                                                     t=f"{now:%H:%M}", a=f"{p_start:%d/%m/%Y}", b=f"{p_end:%d/%m/%Y}"))
    ws.column_dimensions["A"].width = 44
    for col in "BCDEF":
        ws.column_dimensions[col].width = 16
    ws["A4"] = tr("Indicateurs clés")
    ws["A4"].font = F_SECTION
    kpis = [
        (tr("Jours de la période"), f"=COUNTA({Q.rng('date')})", FMT_INT),
        (tr("Jours renseignés"), f'=COUNTIFS({Q.rng("logged")},"{YES}")', FMT_INT),
        (tr("Régularité"), "=IFERROR(B6/B5,0)", FMT_PCT),
        (tr("Intentions fixées"), f"=COUNTA({I.rng('id')})", FMT_INT),
        (tr("Intentions évaluées"), f'=COUNTIFS({I.rng("status")},"<>{S_PREVU}",{I.rng("id")},"<>")', FMT_INT),
        (tr("Intentions réalisées (totalement ou en partie)"),
         f'=COUNTIFS({I.rng("status")},"{S_REAL}")+COUNTIFS({I.rng("status")},"{S_PART}")', FMT_INT),
        (tr("Taux de réalisation"), f'=IFERROR(SUM({I.rng("score")})/B9,"")', FMT_PCT),
        (tr("Temps total suivi (heures)"), f"=SUM({A.rng('min')})/60", FMT_DEC),
        (tr("Temps moyen suivi par jour renseigné (minutes)"), f"=IFERROR(SUM({A.rng('min')})/B6,\"\")", FMT_DEC),
        (tr("Écart moyen entre temps prévu et temps réel (minutes)"), f'=IFERROR(AVERAGE({I.rng("gap")}),"")', FMT_DEC),
        (tr("Sommeil moyen (heures)"), f'=IFERROR(AVERAGE({Q.rng("sleep_hours")}),"")', FMT_DEC),
        (tr("Énergie moyenne (sur 5)"), f'=IFERROR(AVERAGE({Q.rng("energy")}),"")', "0.00"),
        (tr("Humeur moyenne (sur 5)"), f'=IFERROR(AVERAGE({Q.rng("mood")}),"")', "0.00"),
        (tr("Réflexions écrites"), f"=COUNTA({R.rng('content')})", FMT_INT),
        (tr("Plus grand nombre d'heures en une semaine"), f"=MAX({tw.rng('hours')})", FMT_DEC),
        (tr("Journées renseignées le jour même"), f'=COUNTIFS({Q.rng("entry")},"{E_ON}")', FMT_INT),
        (tr("Journées rattrapées plus tard"), f'=COUNTIFS({Q.rng("entry")},"{E_LATE}")', FMT_INT),
        (tr("Journées oubliées (délai dépassé)"), f'=COUNTIFS({Q.rng("entry")},"{E_MISSED}")', FMT_INT),
        (tr("Taux d'oubli le jour même"), "=IFERROR((B21+B22)/(B20+B21+B22),\"\")", FMT_PCT),
        (tr("Délai moyen de rattrapage (jours)"),
         f'=IFERROR(AVERAGEIFS({Q.rng("entry_delay")},{Q.rng("entry")},"{E_LATE}"),"")', FMT_DEC),
    ]
    for k, (label, formula, fmt) in enumerate(kpis):
        r = 5 + k
        ws.cell(row=r, column=1, value=label).font = F_BODY
        c = ws.cell(row=r, column=2, value=formula)
        c.font, c.number_format, c.fill = F_BOLD, fmt, FILL_KPI
        ws.cell(row=r, column=1).border = ws.cell(row=r, column=2).border = BORDER_ROW
    ws.cell(row=11, column=3, value=tr("Réalisé = 1, Partiel = 0,5, autres statuts = 0")).font = F_SUB

    r0 = 5 + len(kpis) + 2
    ws.cell(row=r0, column=1, value=tr("Par dimension de vie")).font = F_SECTION
    td_ = Table(ws, [Col("dim", tr("Dimension"), 44), Col("hours", tr("Heures"), 16, FMT_DEC), Col("n", tr("Intentions"), 16, FMT_INT),
                     Col("rate", tr("Taux de réalisation"), 16, FMT_PCT), Col("share", tr("Part du temps"), 16, FMT_PCT)],
                header_row=r0 + 1)
    for lbl in [tr(x) for x in DIM_LABELS]:
        td_.add({
            "dim": lbl,
            "hours": lambda r: f"=SUMIFS({A.rng('min')},{A.rng('dim')},$A{r})/60",
            "n": lambda r: f"=COUNTIFS({I.rng('dim')},$A{r})",
            "rate": lambda r: (f'=IFERROR(SUMIFS({I.rng("score")},{I.rng("dim")},$A{r})/'
                               f'COUNTIFS({I.rng("dim")},$A{r},{I.rng("status")},"<>{S_PREVU}"),"")'),
            "share": lambda r: f'=IFERROR(B{r}/$B$12,"")',
        })
    ws.freeze_panes = None

    # Graphiques
    bar = BarChart()
    bar.type = "bar"
    bar.title = tr("Heures par dimension")
    bar.style = 10
    bar.y_axis.title = tr("heures")
    bar.add_data(Reference(ws, min_col=2, min_row=td_.header_row, max_row=td_.last), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=td_.first, max_row=td_.last))
    bar.legend = None
    bar.series[0].graphicalProperties.solidFill = GREEN
    bar.height, bar.width = 8, 15
    ws.add_chart(bar, "G4")

    line = LineChart()
    line.title = tr("Taux de réalisation par semaine")
    line.style = 12
    line.y_axis.number_format = "0%"
    line.y_axis.scaling.min, line.y_axis.scaling.max = 0, 1
    line.add_data(Reference(ws_w, min_col=8, min_row=tw.header_row, max_row=tw.last), titles_from_data=True)
    line.set_categories(Reference(ws_w, min_col=1, min_row=tw.first, max_row=tw.last))
    line.legend = None
    line.series[0].graphicalProperties.line.solidFill = RED
    line.series[0].graphicalProperties.line.width = 28000
    line.height, line.width = 8, 15
    ws.add_chart(line, "G21")

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    guide_row = td_.last + 3
    ws.cell(row=guide_row, column=1, value=tr("Contenu du classeur")).font = F_SECTION
    guide = [
        (", ".join(tr(x) for x in "Quotidien".split(", ")), tr("Une ligne par jour : indicateurs déclarés et totaux calculés par formules.")),
        (", ".join(tr(x) for x in "Hebdomadaire, Mensuel".split(", ")), tr("Agrégats par semaine et par mois (formules).")),
        (", ".join(tr(x) for x in "Intentions, Activites, Reflexions".split(", ")), tr("Données brutes de la période. Modifiables : les indicateurs se recalculent.")),
        (", ".join(tr(x) for x in "Calibration, Abandons".split(", ")), tr("Temps prévu contre temps réel, objectifs non réalisés et raisons (formules).")),
        (", ".join(tr(x) for x in "Objectifs, Competences, Evaluations".split(", ")), tr("Progression des objectifs et des compétences.")),
        (", ".join(tr(x) for x in "Analyses".split(", ")), tr("Tendances, associations, jours atypiques, recommandations (valeurs calculées par PEI).")),
        (", ".join(tr(x) for x in "Decisions, Experiences, Relations, Medias".split(", ")), tr("Journal des décisions, expériences, relations et fichiers.")),
    ]
    for k, (sheet, txt) in enumerate(guide, start=1):
        ws.cell(row=guide_row + k, column=1, value=sheet).font = F_BOLD
        ws.cell(row=guide_row + k, column=2, value=txt).font = F_BODY

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = sheet is not ws_syn
    ws_syn.sheet_properties.tabColor = GREEN
    for sh in (ws_q, ws_w, ws_mo):
        sh.sheet_properties.tabColor = YELLOW
    ws_an.sheet_properties.tabColor = RED
    return wb


def export_bytes(session: Session, start: date | None = None, end: date | None = None) -> bytes:
    buf = io.BytesIO()
    build_workbook(session, start, end).save(buf)
    return buf.getvalue()


def filename(start: date | None, end: date | None) -> str:
    stamp = datetime.now().strftime("%Y%m%d")
    if start or end:
        return f"PEI_export_{(start or date.min):%Y%m%d}_{(end or date.today()):%Y%m%d}.xlsx"
    return f"{tr('PEI_export_complet')}_{stamp}.xlsx"
