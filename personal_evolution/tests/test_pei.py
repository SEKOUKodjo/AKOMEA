from __future__ import annotations

import io
from datetime import date, timedelta

import pytest
from fastapi.testclient import TestClient

from ai.nlp.extract import extract_intentions, extract_review, parse_minutes


@pytest.fixture()
def client(tmp_path, monkeypatch):
    from backend import config, db

    config.configure(data_dir=tmp_path, storage_dir=tmp_path / "storage", backup_dir=tmp_path / "backups",
                     db_path=tmp_path / "test.db", pin=None)
    db.init_engine()
    from backend.main import app

    with TestClient(app) as c:
        yield c


# Extraction en langage naturel

@pytest.mark.parametrize("text,expected", [
    ("1h30", 90), ("une heure quinze", 75), ("deux heures", 120), ("45 min", 45),
    ("une demi heure", 30), ("trois quarts d'heure", 45), ("une heure et demie", 90),
])
def test_parse_minutes(text, expected):
    assert parse_minutes(text) == expected


def test_morning_example_from_specification():
    r = extract_intentions("Aujourd'hui, je veux lire un chapitre de la Bible, faire 30 minutes de sport, "
                           "travailler deux heures sur mon projet Python et lire un article scientifique.")
    got = [(i.dimension, i.category, i.minutes) for i in r.items]
    assert got == [("spirituel", "Lecture biblique", None), ("physique", "Sport", 30),
                   ("etudes", "Python", 120), ("etudes", "Article scientifique", None)]


def test_priority_markers():
    r = extract_intentions("Je dois absolument envoyer ma candidature et appeler maman si j'ai le temps")
    assert [i.priority for i in r.items] == [1, 3]


def test_evening_example_from_specification():
    intentions = [
        {"id": 1, "description": "Faire 30 minutes de sport", "category": "Sport", "dimension": "physique", "estimated_minutes": 30},
        {"id": 2, "description": "Travailler deux heures sur Python", "category": "Python", "dimension": "etudes", "estimated_minutes": 120},
    ]
    r = extract_review("J'avais prévu deux heures de Python mais j'ai finalement travaillé une heure quinze. "
                       "J'ai terminé la partie sur les modèles de classification. "
                       "Je n'ai pas fait de sport parce que j'étais fatigué.", intentions)
    python = next(i for i in r.items if i.intention_id == 2)
    assert (python.planned_minutes, python.minutes, python.status) == (120, 75, "partiel")
    assert "classification" in python.result
    sport = next(i for i in r.items if i.intention_id == 1)
    assert sport.status == "non_realise" and sport.reason == "j'étais fatigué"
    assert any(e["emotion"] == "fatigue" for e in r.emotions)


def test_activity_with_result():
    r = extract_review("J'ai travaillé 1h30 sur Python et terminé les exercices de régression.")
    item = r.items[0]
    assert (item.dimension, item.category, item.minutes) == ("etudes", "Python", 90)
    assert "régression" in item.result


# Cycle quotidien par l'API

def test_morning_evening_cycle_and_provenance(client):
    today = str(date.today())
    parsed = client.post("/api/morning/parse", json={"text": "Je veux faire 30 minutes de sport et lire un article", "date": today}).json()
    assert len(parsed["items"]) == 2
    items = [{"description": i["description"], "dimension": i["dimension"], "category": i["category"],
              "estimated_minutes": i["minutes"]} for i in parsed["items"]]
    items[1]["estimated_minutes"] = 40  # correction par l'utilisateur
    day = client.post("/api/morning", json={"date": today, "intentions": items, "extraction_id": parsed["extraction_id"],
                                            "metrics": {"sleep_hours": 7.5, "energy": 4}}).json()
    assert day["summary"]["n_intentions"] == 2 and day["day"]["sleep_hours"] == 7.5

    src = client.get(f"/api/sources/{parsed['source_id']}").json()
    assert src["raw_text"].startswith("Je veux")
    assert src["extractions"][0]["status"] == "corrige"

    sport, article = day["intentions"]
    out = client.post("/api/evening", json={
        "date": today,
        "reviews": [{"intention_id": sport["id"], "status": "non_realise", "reason": "pluie"},
                    {"intention_id": article["id"], "status": "realise", "actual_minutes": 55, "result": "Article résumé"}],
        "activities": [{"description": "Appel à maman", "dimension": "relationnel", "duration_minutes": 20}],
        "reflections": [{"content": "J'ai compris la validation croisée", "kind": "apprentissage"}],
        "metrics": {"mood": 4, "highlight": "Belle journée"},
    }).json()
    assert out["summary"]["completion_rate"] == 0.5
    assert out["summary"]["actual_minutes"] == 75
    assert out["day"]["evening_done_at"]

    search = client.get("/api/memory/search", params={"q": "validation croisée"}).json()
    assert search and search[0]["kind"] == "reflexion"
    cal = client.get("/api/statistics/calibration").json()
    assert cal["n"] == 1 and cal["overall"]["mean_error"] == 15


def test_activity_updates_intention(client):
    it = client.post("/api/intentions", json={"description": "Python", "category": "Python", "estimated_minutes": 60}).json()
    client.post("/api/activities", json={"description": "Python", "intention_id": it["id"], "start_time": "09:00", "end_time": "09:40"})
    day = client.get(f"/api/days/{date.today()}").json()
    assert day["intentions"][0]["status"] == "partiel" and day["intentions"][0]["actual_minutes"] == 40


def test_audio_upload_keeps_original(client):
    r = client.post("/api/audio", files={"file": ("matin.m4a", io.BytesIO(b"fake audio"), "audio/mp4")},
                    data={"moment": "matin"})
    media = r.json()
    assert media["type"] == "audio" and media["sha256"] and media["transcription_status"] in {"en_attente", "indisponible"}
    assert client.get(f"/api/media/{media['id']}/file").content == b"fake audio"


def test_goals_hierarchy_and_skills(client):
    vision = client.post("/api/goals", json={"title": "Vision", "horizon": "vision"}).json()
    child = client.post("/api/goals", json={"title": "Projet", "horizon": "annuel", "parent_id": vision["id"]}).json()
    client.post("/api/intentions", json={"description": "Avancer", "goal_id": child["id"]})
    prog = {g["id"]: g for g in client.get("/api/goals/progress").json()}
    assert prog[vision["id"]]["n_intentions"] == 1  # remonte vers la vision
    sk = client.post("/api/skills", json={"name": "Pandas"}).json()
    client.post("/api/activities", json={"description": "Exercices Pandas", "duration_minutes": 30})
    client.post(f"/api/skills/{sk['id']}/progress", json={"self_rating": 3})
    skills = client.get("/api/skills").json()
    assert skills[0]["hours"] == 0.5 and skills[0]["current_rating"] == 3


def test_statistics_on_demo_data(client):
    from backend.db import session_scope
    from backend.services import demo

    with session_scope() as s:
        demo.seed(s, n_days=90)
    summary = client.get("/api/statistics/summary", params={"days": 30}).json()
    assert summary["n_days"] == 30 and 0 <= summary["completion_rate"] <= 1
    assert client.get("/api/statistics/correlations").json()["note"]
    assert isinstance(client.get("/api/insights").json(), list)
    pred = client.post("/api/predict/completion", json={"category": "Sport", "dimension": "physique", "estimated_minutes": 30}).json()
    assert 0 <= pred["probability"] <= 1 and pred["disclaimer"]
    sim = client.post("/api/simulate", json={"category": "Python", "minutes_per_day": 90, "days": 90}).json()
    assert sim["hours_p10"] <= sim["hours_median"] <= sim["hours_p90"] <= sim["ideal_hours"] * 3
    ans = client.get("/api/memory/ask", params={"q": "Quels sujets ai-je le plus étudiés ?"}).json()
    assert ans["intent"] == "temps" and ans["items"]
    exps = client.get("/api/experiments").json()
    assert exps and "message" in exps[0]["analysis"]
    yesterday = date.today() - timedelta(days=1)
    assert isinstance(client.get(f"/api/predict/day/{yesterday}").json(), list)


def test_backup(client):
    client.post("/api/reflections", json={"content": "Note"})
    r = client.post("/api/system/backup").json()
    assert r["file"].endswith(".zip")


def test_pin_protection(client):
    from backend import config

    config.settings.pin = "1234"
    try:
        assert client.get("/api/goals").status_code == 401
        assert client.get("/api/goals", headers={"X-PEI-PIN": "1234"}).status_code == 200
    finally:
        config.settings.pin = None


def test_frontend_and_icons(client):
    html = client.get("/").text
    assert 'id="i-sun"' in html  # sprite d'icônes intégré
    assert client.get("/icons/icon-192.png").content[:4] == b"\x89PNG"
    assert client.get("/manifest.json").json()["theme_color"] == "#006A4E"


def test_excel_export(client):
    from openpyxl import load_workbook

    empty = client.get("/api/export/excel")
    assert empty.status_code == 200 and empty.content[:2] == b"PK"  # base vide : classeur valide

    from backend.db import session_scope
    from backend.services import demo

    with session_scope() as s:
        demo.seed(s, n_days=40)
    r = client.get("/api/export/excel", params={"days": 30})
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"] and ".xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    for sheet in ["Synthese", "Quotidien", "Hebdomadaire", "Mensuel", "Intentions", "Activites", "Calibration",
                  "Abandons", "Objectifs", "Competences", "Analyses", "Decisions"]:
        assert sheet in wb.sheetnames
    q = wb["Quotidien"]
    assert q.max_row == 31  # en-tête + 30 jours
    assert str(q["D2"].value).startswith("=COUNTIFS(")  # indicateurs calculés par formules
    assert wb["Synthese"]["B11"].value.startswith("=IFERROR(SUM(")
