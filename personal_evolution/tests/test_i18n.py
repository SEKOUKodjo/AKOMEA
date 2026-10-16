"""Bilinguisme : chaque texte traduit doit exister dans les dictionnaires anglais."""
from __future__ import annotations

import ast
import io
import json
import re
import subprocess
from pathlib import Path

import pytest

from backend import i18n
from backend.locales.en import EN

ROOT = Path(__file__).resolve().parent.parent
PY_DIRS = ["backend", "ai", "dashboard"]
JS_APP = ROOT / "frontend" / "src" / "app.js"
JS_EN = ROOT / "frontend" / "src" / "i18n_en.js"
TEMPLATE_VARS = {"template", "text", "msg", "message"}


def _consts(node) -> list[str]:
    """Chaînes possibles d'une expression : constante, ou branches d'une condition."""
    if isinstance(node, ast.Constant) and isinstance(node.value, str):
        return [node.value]
    if isinstance(node, ast.IfExp):
        return _consts(node.body) + _consts(node.orelse)
    if isinstance(node, ast.Dict):
        return [s for v in node.values for s in _consts(v)]
    return []


def python_keys() -> set[str]:
    """Textes passés à tr() dans le code Python, et modèles de phrases préparés avant l'appel."""
    keys: set[str] = set()
    for folder in PY_DIRS:
        for path in (ROOT / folder).rglob("*.py"):
            if "locales" in path.parts:
                continue
            tree = ast.parse(path.read_text(encoding="utf-8"))
            for node in ast.walk(tree):
                if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id == "tr" and node.args:
                    keys.update(_consts(node.args[0]))
                    arg = node.args[0]
                    if isinstance(arg, ast.Subscript) and isinstance(arg.value, ast.Dict):
                        keys.update(_consts(arg.value))
                elif isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id in TEMPLATE_VARS
                                                          for t in node.targets):
                    keys.update(s for s in _consts(node.value) if re.search(r"[a-zà-ÿ]{3}", s) and " " in s)
    return keys


def dynamic_keys() -> set[str]:
    """Libellés de référence traduits à l'exécution (dimensions, statuts, indicateurs...)."""
    from ai.ml import predict
    from backend import domain
    from backend.services import analytics, memory

    keys: set[str] = set(i18n.WEEKDAYS) | set(i18n.WEEKDAYS_SHORT)
    keys |= {v["label"] for v in domain.DIMENSIONS.values()}
    for mapping in (domain.HORIZONS, domain.GOAL_STATUSES, domain.INTENTION_STATUSES, domain.REFLECTION_KINDS,
                    domain.MOMENTS, analytics.ENTRY_STATUS, memory.KIND_LABELS):
        keys |= set(mapping.values())
    keys |= {v if isinstance(v, str) else v.get("label", "") for v in domain.DAY_METRICS.values()}
    keys |= set(analytics.METRIC_LABELS.values())
    keys |= {analytics.CAUSALITY_NOTE, predict.DISCLAIMER, "Bilan de la journée", "Saisie matin", "Saisie journee",
             "Saisie soir", "Saisie"}
    keys |= {"sous estimation", "surestimation", "juste", "insuffisant", "hausse", "baisse", "haut", "bas",
             "1 à 3", "4", "5", "6 et plus"}
    keys |= {"Activité", "Compétence", "Décision", "Expérience", "Fichier", "Intention", "Média", "Objectif",
             "Personne", "Réflexion", "Source", "Élément"}
    keys |= {"exercice", "projet", "test", "formation", "production"}
    return {k for k in keys if k and not k.isdigit()}


# Textes que app.js traduit à travers une variable (statuts, échelles, tuiles, suggestions...).
JS_DYNAMIC = {
    "Réalisé", "Partiel", "Non réalisé", "Abandonné", "Reporté", "Prévu",
    "Bonjour", "Bon après midi", "Bonsoir", "Accueil", "Matin", "Journée", "Soir", "Plus",
    "faible", "élevé", "mauvaise", "excellente", "basse", "calme", "très stressé", "débutant", "expert",
    "Objectifs", "Vision et objectifs", "Compétences", "Ressenti et preuves", "Journal", "Décisions et expériences",
    "Relations", "Personnes à entretenir", "Mémoire", "Retrouver et interroger", "Bilan", "Tendances et prédictions",
    "Historique", "Toutes les journées", "Tableau de bord", "Analyse complète sur le PC", "Export Excel",
    "Données et indicateurs", "Réglages", "Langue, sauvegarde et réseau",
    "appel", "visite", "message", "rencontre", "exercice", "projet", "test", "formation", "production",
    "famille", "ami", "collègue", "mentor", "communauté",
    "Qu'est ce qui m'a marqué le mois dernier ?", "Quels projets ai je terminés cette année ?",
    "Quels sujets ai je le plus étudiés ?", "Quelles difficultés reviennent régulièrement ?",
    "Comment ai je évolué ce mois ci ?",
}


def js_keys() -> set[str]:
    src = JS_APP.read_text(encoding="utf-8")
    keys = set(JS_DYNAMIC)
    for match in re.finditer(r'\btr\(\s*"((?:[^"\\]|\\.)*)"', src):
        keys.add(json.loads(f'"{match.group(1)}"'))
    for match in re.finditer(r"\btr\(\s*'((?:[^'\\]|\\.)*)'", src):
        keys.add(match.group(1).replace("\\'", "'"))
    for match in re.finditer(r'data-t>([^<]+)<', (ROOT / "frontend" / "index.html").read_text(encoding="utf-8")):
        keys.add(match.group(1).strip())
    return keys


def js_dictionary() -> dict[str, str]:
    """Lit PEI_EN avec Node si disponible, sinon par expression régulière."""
    src = JS_EN.read_text(encoding="utf-8")
    try:
        out = subprocess.run(["node", "-e", src + "\nprocess.stdout.write(JSON.stringify(PEI_EN));"],
                             capture_output=True, text=True, timeout=30, check=True)
        return json.loads(out.stdout)
    except (FileNotFoundError, subprocess.CalledProcessError):
        return {json.loads(f'"{k}"'): v for k, v in re.findall(r'"((?:[^"\\]|\\.)*)"\s*:\s*"((?:[^"\\]|\\.)*)"', src)}


def _placeholders(s: str) -> set[str]:
    return set(re.findall(r"\{(\w+)\}", s))


def test_python_keys_translated():
    missing = sorted((python_keys() | dynamic_keys()) - set(EN))
    assert not missing, f"{len(missing)} textes sans traduction anglaise : {missing[:20]}"


def test_python_placeholders_match():
    wrong = [k for k, v in EN.items() if _placeholders(k) != _placeholders(v)]
    assert not wrong, wrong[:10]


def test_js_keys_translated():
    dico = js_dictionary()
    missing = sorted(js_keys() - set(dico))
    assert not missing, f"{len(missing)} textes sans traduction anglaise : {missing[:20]}"
    wrong = [k for k, v in dico.items() if _placeholders(k) != _placeholders(v)]
    assert not wrong, wrong[:10]


def test_tr_and_numbers():
    with i18n.use("en"):
        assert i18n.tr("Objectif") == "Goal"
        assert i18n.num(2.5) == "2.5"
        assert i18n.weekday(0) == "Monday"
    with i18n.use("fr"):
        assert i18n.tr("Objectif") == "Objectif"
        assert i18n.num(2.5) == "2,5"
    assert i18n.from_accept_language("en-GB,en;q=0.9,fr;q=0.8") == "en"
    assert i18n.normalize("de") == "fr"


@pytest.mark.parametrize("text,dim,cat,minutes", [
    ("Read the Bible for 30 minutes", "spirituel", "Bible reading", 30),
    ("Study machine learning for 2 hours", "etudes", "Machine Learning", 120),
    ("Go running for 45 min", "physique", "Running", 45),
    ("Call my mom", "relationnel", "Family", None),
])
def test_english_extraction(text, dim, cat, minutes):
    from ai.nlp.extract import extract_intentions

    res = extract_intentions(text, out_lang="en")
    assert res.language == "en"
    it = res.items[0]
    assert it.dimension == dim and it.category == cat
    assert it.minutes == minutes


# API en anglais

@pytest.fixture()
def client(tmp_path):
    from fastapi.testclient import TestClient

    from backend import config, db

    config.configure(data_dir=tmp_path, storage_dir=tmp_path / "storage", backup_dir=tmp_path / "backups",
                     db_path=tmp_path / "test.db", backfill_days=4)
    db.init_engine()
    from backend.main import app

    with TestClient(app) as c:
        yield c


def test_language_resolution_and_preference(client):
    assert client.get("/api/auth/status", headers={"Accept-Language": "en-US,en;q=0.9"}).json()["language"] == "en"
    assert client.get("/api/auth/status", headers={"Accept-Language": "fr-FR"}).json()["language"] == "fr"
    r = client.get("/api/goals", headers={"X-PEI-Lang": "en"})
    assert r.status_code == 401 and r.json()["detail"] == "Sign in required"
    assert client.get("/api/goals").json()["detail"] == "Connexion requise"

    r = client.post("/api/auth/setup", json={"username": "kodjo", "password": "Togo2026x"}, headers={"X-PEI-Lang": "en"})
    assert r.json()["user"]["language"] == "en"
    r = client.patch("/api/auth/preferences", json={"language": "fr"})
    assert r.json()["user"]["language"] == "fr" and client.cookies.get(i18n.COOKIE) == "fr"
    assert client.patch("/api/auth/preferences", json={"language": "de"}).status_code == 422
    # Le cookie de langue suffit ensuite, sans en tête.
    client.cookies.set(i18n.COOKIE, "en")
    meta = client.get("/api/meta").json()
    assert meta["lang"] == "en"
    assert {d["label"] for d in meta["dimensions"]} >= {"Spiritual", "Studies", "Leisure and creativity"}
    assert {s["label"] for s in meta["statuses"]} >= {"Done", "Not done"}
    assert client.get("/api/goals/999999", headers={"X-PEI-Lang": "en"}).status_code in (404, 405)


def test_english_morning_and_statistics(client):
    from datetime import date

    from backend.db import session_scope
    from backend.services import demo

    client.post("/api/auth/setup", json={"username": "kodjo", "password": "Togo2026x"})
    with session_scope() as s:
        demo.seed(s, n_days=40)
    h = {"X-PEI-Lang": "en"}
    parsed = client.post("/api/morning/parse", headers=h,
                         json={"text": "Today I want to pray, run for 30 minutes and study Python for two hours",
                               "date": str(date.today())}).json()
    assert [(i["dimension"], i["category"], i["minutes"]) for i in parsed["items"]] == [
        ("spirituel", "Prayer", None), ("physique", "Running", 30), ("etudes", "Python", 120)]

    stats = client.get("/api/statistics/summary", headers=h, params={"days": 30})
    assert stats.status_code == 200
    assert {d["label"] for d in stats.json()["dimensions"]} >= {"Spiritual", "Studies"}
    insights = client.get("/api/insights", headers=h)
    assert insights.status_code == 200 and insights.json()
    french = [i["message"] for i in insights.json() if re.search(r"\b(les|des|tes|sont|avec|pour)\b", i["message"])]
    assert not french, french
    answer = client.get("/api/memory/ask", headers=h, params={"q": "What did I study the most last month?"})
    assert answer.status_code == 200
    assert answer.json()["period"] == "last month", answer.json()


def test_english_excel_export(client):
    from openpyxl import load_workbook

    from backend.db import session_scope
    from backend.services import demo

    client.post("/api/auth/setup", json={"username": "kodjo", "password": "Togo2026x"})
    with session_scope() as s:
        demo.seed(s, n_days=40)
    r = client.get("/api/export/excel", params={"days": 30}, headers={"X-PEI-Lang": "en"})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    for sheet in ["Summary", "Daily", "Weekly", "Monthly", "Intentions", "Activities", "Calibration", "Skills"]:
        assert sheet in wb.sheetnames
    assert str(wb["Daily"]["D2"].value).startswith("=COUNTIFS(")
